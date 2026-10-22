"""Compile the measured deck spec into one scene and export native PowerPoint.

PPT Master owns SVG geometry conversion. python-pptx owns one PPTX package,
text, pictures, charts and tables. No image generation or raster tracing occurs here.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import sys
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SVG_NS = 'http://www.w3.org/2000/svg'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
EMU = 9525
FONTS = {'zh': 'Microsoft YaHei', 'en': 'Times New Roman'}
CORE_FIELDS = ('id','kind','box','text','label','runs','paragraph','line_breaks','font_size_pt',
               'bold','italic','color','alignment','geometry','visual_style','corner_radius_px',
               'svg_path','svg_asset_path','asset_path','role','layer','source_ids',
               'from','to','start_point','end_point','chart_type','categories','series','unit',
               'rows','component_id','registration_group','layer_role','master_slot','adjustments')


def digest(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write(path, value):
    path=Path(path)
    if path.exists():raise ValueError('输出已存在，请使用新版本文件名: '+str(path))
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def finite(value):
    return isinstance(value,(float,int)) and not isinstance(value,bool) and math.isfinite(value)


def asset(base, relative):
    path=(Path(base)/relative).resolve()
    if not path.is_file():raise ValueError('素材不存在: '+str(path))
    return path


def compile_scene(spec, base):
    """Visible content/geometry only; hidden notes and production records are excluded."""
    if spec.get('mode')!='per-slide':raise ValueError('保持逐页生图，不新增制作路线')
    if spec.get('master',{}).get('tokens',{}).get('fonts')!=FONTS:
        raise ValueError('保留微软雅黑与 Times New Roman')
    scene={'scene_version':1,'canvas':copy.deepcopy(spec['canvas']),
           'fonts':FONTS.copy(),'master':{'elements':[]},'slides':[]}
    def element_copy(element):
        style=spec['master'].get('tokens',{}).get('text_styles',{}).get(element.get('master_slot'),{})
        element={**style,**element}
        out={k:copy.deepcopy(element[k]) for k in CORE_FIELDS if k in element}
        for key in ('asset_path','svg_asset_path'):
            if out.get(key):out[key.replace('_path','_sha256')]=hashlib.sha256(asset(base,out[key]).read_bytes()).hexdigest()
        return out
    scene['master']['elements']=[element_copy(e) for e in spec['master'].get('elements',[])]
    for page in spec['slides']:
        scene['slides'].append({'id':page['id'],'background':page.get('background','#FFFFFF'),
            'elements':[element_copy(e) for e in page['elements']]})
    errors=validate_scene(scene,base)
    if errors:raise ValueError('；'.join(errors))
    return scene


def validate_scene(scene, base):
    from PIL import Image
    errors=[]
    if scene.get('scene_version')!=1:errors.append('scene_version 无效')
    canvas=scene.get('canvas',{})
    if not all(finite(canvas.get(k)) and canvas[k]>0 for k in ('width','height')):return ['canvas 无效']
    if scene.get('fonts')!=FONTS:errors.append('字体设置已改变')
    pages=scene.get('slides',[])
    if not pages:errors.append('没有页面')
    page_ids=[p.get('id') for p in pages]
    if len(set(page_ids))!=len(page_ids):errors.append('页面 id 重复')
    groups=[('master',scene.get('master',{}).get('elements',[]))]+[(p.get('id'),p.get('elements',[])) for p in pages]
    for sid,elements in groups:
        ids=[e.get('id') for e in elements]
        if len(set(ids))!=len(ids):errors.append(str(sid)+': 对象 id 重复')
        for e in elements:
            label=f"{sid}/{e.get('id')}"
            if not e.get('id'):errors.append(label+': 缺少 id')
            b=e.get('box',{})
            if not all(finite(b.get(k)) for k in ('left','top','width','height')):
                errors.append(label+': box 无效');continue
            if min(b.values())<0 or b['left']+b['width']>canvas['width']+.1 or b['top']+b['height']>canvas['height']+.1:
                errors.append(label+': box 越界')
            if sid=='master' and e.get('kind') not in {'text','shape','image'}:errors.append(label+': 固定母版对象仅支持文字、图形和图片')
            if e.get('kind')=='text':
                runs=e.get('runs',[])
                if runs and ''.join(r.get('text','') for r in runs).replace('\n','')!=e.get('text','').replace('\n',''):
                    errors.append(label+': runs 与准确文字不一致')
                for run in runs:
                    if run.get('font_family') not in FONTS.values():errors.append(label+': runs 使用非指定字体')
            if e.get('kind')=='image':
                try:
                    path=asset(base,e['asset_path'])
                    if hashlib.sha256(path.read_bytes()).hexdigest()!=e.get('asset_sha256'):errors.append(label+': 图片已改变')
                    with Image.open(path) as image:w,h=image.size
                    if not b['height'] or abs((b['width']/b['height'])/(w/h)-1)>.005:
                        errors.append(label+': 先按原图比例调整 spec.box，不能拉伸图片')
                except (OSError,KeyError,ValueError) as error:errors.append(label+': '+str(error))
            if e.get('kind')=='shape' and e.get('svg_asset_path'):
                try:
                    path=asset(base,e['svg_asset_path'])
                    if hashlib.sha256(path.read_bytes()).hexdigest()!=e.get('svg_asset_sha256'):errors.append(label+': 矢量素材已改变')
                    safe_vector(path.read_bytes())
                except (OSError,ValueError,ET.ParseError) as error:errors.append(label+': '+str(error))
            if e.get('kind')=='connector' and (e.get('from') not in ids or e.get('to') not in ids):errors.append(label+': 连接端点不存在')
    # Full-canvas registered image layers must share placement; do not trim them.
    for page in pages:
        registered={}
        for e in page['elements']:
            if e.get('registration_group'):
                group=registered.setdefault(e['registration_group'],[]);group.append(e)
        for gid,members in registered.items():
            if any(e['box']!=members[0]['box'] for e in members):errors.append(page['id']+'/'+gid+': 注册图层画布或位置不一致')
    return errors


def safe_vector(data):
    if len(data)>2_000_000 or re.search(br'<!DOCTYPE|<!ENTITY',data,re.I):raise ValueError('矢量素材包含不支持的 XML 声明或过大')
    tree=ET.fromstring(data)
    allowed={'svg','g','path','rect','circle','ellipse','line','polyline','polygon','defs','linearGradient','radialGradient','stop'}
    for node in tree.iter():
        if node.tag.rsplit('}',1)[-1] not in allowed:raise ValueError('首批矢量接入仅支持图形，不接受内嵌文字、图片或脚本')
        if any(key.rsplit('}',1)[-1] in {'href','onload','onclick'} for key in node.attrib):raise ValueError('矢量素材不能引用外部资源或事件')
    return tree


def load_converter():
    folder=ROOT/'vendor/ppt-master'
    manifest=load(folder/'vendor-manifest.json')
    for file in manifest['files']:
        path=folder/file['path']
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=file['sha256']:
            raise ValueError('PPT Master 文件缺失或改变: '+file['path'])
    path=str(folder/'scripts')
    if path not in sys.path:sys.path.insert(0,path)
    from attribution_guard import require_skill_integrity
    require_skill_integrity()
    from svg_to_pptx import convert_svg_to_slide_shapes
    return convert_svg_to_slide_shapes


def color(value, default='#000000'):
    value=value or default
    if value=='none':return 'none'
    if not re.fullmatch('#[0-9A-Fa-f]{6}',value):raise ValueError('颜色需要 #RRGGBB: '+str(value))
    return value


def shape_svg(e, canvas, base):
    ET.register_namespace('',SVG_NS)
    root=ET.Element('{'+SVG_NS+'}svg',{'viewBox':f"0 0 {canvas['width']} {canvas['height']}"})
    b=e['box'];x,y,w,h=(b[k] for k in ('left','top','width','height'))
    style=e.get('visual_style',{})
    attrs={'id':e['id'],'fill':'none' if style.get('fill_none') else color(style.get('fill'),'#FFFFFF'),
           'stroke':color(style.get('stroke'),'none'),'stroke-width':str(style.get('stroke_width_pt',0)*4/3)}
    if style.get('dash') in {'dash','dot'}:attrs['stroke-dasharray']='6 4' if style['dash']=='dash' else '1 3'
    geometry=e.get('geometry','rect')
    if e.get('svg_asset_path'):
        node=safe_vector(asset(base,e['svg_asset_path']).read_bytes())
        vb=[float(v) for v in node.get('viewBox','').replace(',',' ').split()]
        if len(vb)!=4 or min(vb[2:])<=0:raise ValueError('矢量素材需要有效 viewBox')
        scale=min(w/vb[2],h/vb[3]);dx=x+(w-vb[2]*scale)/2-vb[0]*scale;dy=y+(h-vb[3]*scale)/2-vb[1]*scale
        group=ET.SubElement(root,'g',{'id':e['id'],'transform':f'translate({dx} {dy}) scale({scale})'})
        for child in node:group.append(copy.deepcopy(child))
    elif geometry in {'rect','rectangle','roundRect','rounded-rectangle'}:
        attrs.update(x=str(x),y=str(y),width=str(w),height=str(h))
        if geometry in {'roundRect','rounded-rectangle'}:attrs['rx']=str(e.get('corner_radius_px',min(w,h)*.08))
        ET.SubElement(root,'rect',attrs)
    elif geometry in {'ellipse','circle'}:
        attrs.update(cx=str(x+w/2),cy=str(y+h/2),rx=str(w/2),ry=str(h/2));ET.SubElement(root,'ellipse',attrs)
    elif geometry in {'triangle','chevron','hexagon'}:
        points={'triangle':[(x+w/2,y),(x+w,y+h),(x,y+h)],
                'chevron':[(x,y),(x+w*.75,y),(x+w,y+h/2),(x+w*.75,y+h),(x,y+h),(x+w*.25,y+h/2)],
                'hexagon':[(x+w*.25,y),(x+w*.75,y),(x+w,y+h/2),(x+w*.75,y+h),(x+w*.25,y+h),(x,y+h/2)]}[geometry]
        attrs['points']=' '.join(f'{a},{b}' for a,b in points);ET.SubElement(root,'polygon',attrs)
    elif geometry.startswith('preset:'):
        from pptx_to_svg.preset_authoring import render_preset_shape_fragment
        fragment=render_preset_shape_fragment(geometry.split(':',1)[1],(x,y,w,h),
            element_id=e['id'],adjustments=e.get('adjustments',{}),
            style={k:attrs[k] for k in ('fill','stroke','stroke-width') if k!='stroke-width' or attrs['stroke']!='none'})
        root.append(ET.fromstring(fragment))
    elif geometry in {'freeform','path'} and e.get('svg_path'):
        # Paths use page coordinates, matching the source scene and preview.
        attrs['d']=e['svg_path'];ET.SubElement(root,'path',attrs)
    else:raise ValueError('未支持的原生图形，请使用有依据的 SVG 或独立图片: '+geometry)
    return ET.tostring(root,encoding='unicode')


def put_text(shape, e):
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.dml.color import RGBColor
    frame=shape.text_frame;frame.clear()
    paragraph=e.get('paragraph',{})
    margins=paragraph.get('margins_px',{})
    for side in ('left','right','top','bottom'):setattr(frame,'margin_'+side,int(margins.get(side,0)*EMU))
    frame.word_wrap=True
    frame.vertical_anchor={'top':MSO_ANCHOR.TOP,'center':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}.get(paragraph.get('vertical_alignment'),MSO_ANCHOR.TOP)
    text=e.get('text',e.get('label',''))
    runs=e.get('runs') or [{'text':text,'font_size_pt':e.get('font_size_pt',20),'bold':e.get('bold',False),'italic':e.get('italic',False),'color':e.get('color','#000000')}]
    plain=''.join(r['text'] for r in runs)
    if '\n' in plain:
        boundaries=[];position=0
        for char in plain:
            if char=='\n':boundaries.append(position)
            else:position+=1
    else:boundaries=e.get('line_breaks',[])
    current=frame.paragraphs[0]
    def configure(p):
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(paragraph.get('alignment',e.get('alignment')),PP_ALIGN.LEFT)
        for field in ('space_before_pt','space_after_pt'):
            setattr(p,field.replace('_pt',''),Pt(paragraph.get(field,0)))
        if paragraph.get('line_spacing_pt'):p.line_spacing=Pt(paragraph['line_spacing_pt'])
        if paragraph.get('continuation_indent_px') is not None:p._p.get_or_add_pPr().set('marL',str(int(paragraph['continuation_indent_px']*EMU)))
    configure(current);offset=0
    for source in runs:
        content=source['text'].replace('\n','')
        # Keep explicit run styles; split unstyled mixed-language text by script.
        chunks=[]
        for char in content:
            family=source.get('font_family') or (FONTS['zh'] if '\u3400'<=char<='\u9fff' or '\u3000'<=char<='\u303f' or '\uff00'<=char<='\uffef' else FONTS['en'])
            if chunks and chunks[-1][0]==family and offset not in boundaries:chunks[-1][1]+=char
            else:chunks.append([family,char,offset])
            offset+=1
        for family,chunk,start in chunks:
            if start in boundaries and start:
                current=frame.add_paragraph();configure(current)
            run=current.add_run();run.text=chunk
            run.font.name=family;run.font.size=Pt(source.get('font_size_pt',e.get('font_size_pt',20)))
            run.font.bold=source.get('bold',e.get('bold',False));run.font.italic=source.get('italic',e.get('italic',False))
            run.font.color.rgb=RGBColor.from_string(color(source.get('color',e.get('color'))).lstrip('#'))
            prop=run._r.get_or_add_rPr()
            for tag in ('a:latin','a:ea','a:cs'):
                node=prop.find(tag.replace('a:',f'{{{A}}}'))
                if node is None:node=OxmlElement(tag);prop.append(node)
                node.set('typeface',family)


def layer_jobs(spec, base):
    """Reference-edit requests for independent layers; never regenerate evidence."""
    jobs=[]
    for page in spec['slides']:
        candidates=[e for e in page['elements'] if e.get('kind')=='image' and e.get('layer_role')]
        for e in candidates:
            if e.get('role') in {'evidence','logo'}:continue
            canonical=asset(base,page['design_reference'])
            role=e['layer_role']
            if role not in {'base','subject','foreground','cutout'}:raise ValueError('layer_role 无效')
            text=('根据本次提供的原始设计图，提取或参考编辑指定图层；保持原画布、位置、尺寸与视觉风格。'
                  '仅处理本层，不改变数据、标志或主体身份。不加入文字、数字或标签。')
            if role=='base':text+='制作干净背景，移除已计划独立重建的正文与主体，仅补全露出的背景。'
            else:text+='保留本层主体，其他区域透明；维持原画布，不自动裁成紧边界。'
            jobs.append({'slide_id':page['id'],'element_id':e['id'],'layer_role':role,
                'canonical_reference':{'path':str(canonical),'sha256':hashlib.sha256(canonical.read_bytes()).hexdigest()},
                'reference_box':e.get('reference_box'),'placement':e['box'],'registration_group':e.get('registration_group'),
                'canvas':spec['canvas'],'transparent_background':role!='base',
                'description':e.get('asset_prompt',''),'prompt':text+e.get('asset_prompt','')})
    return {'jobs':jobs,'note':'仅准备内置 imagegen 任务；生成和检查真实素材后再写 asset_path，不伪造图片。'}


def geometry_shape(slide, e, scene, base, scratch, converter):
    from pptx.oxml import parse_xml
    from pptx.oxml.xmlchemy import OxmlElement
    svg=scratch/(hashlib.sha256(e['id'].encode()).hexdigest()[:20]+'.svg')
    svg.write_text(shape_svg(e,scene['canvas'],base),encoding='utf-8')
    events=[]
    xml,media,rels,animations,parts,types=converter(svg,resource_root=Path(base),promote_background=False,trace_out=events,image_optimize=False)
    if media or rels or parts:raise ValueError('几何转换不允许混入图片或外部关系')
    for trace in events:
        if trace.get('summary',{}).get('skipped'):raise ValueError('存在未转换的矢量元素')
    tree=parse_xml(xml.encode('utf-8')).find(f'{{{P}}}cSld/{{{P}}}spTree')
    objects=[copy.deepcopy(n) for n in tree if n.tag.rsplit('}',1)[-1] in {'sp','grpSp','cxnSp'}]
    if not objects:raise ValueError('没有生成原生图形')
    if len(objects)==1:
        node=objects[0]
    else:
        node=OxmlElement('p:grpSp');nv=OxmlElement('p:nvGrpSpPr');nv.append(OxmlElement('p:cNvPr'));nv.append(OxmlElement('p:cNvGrpSpPr'));nv.append(OxmlElement('p:nvPr'));node.append(nv)
        prop=OxmlElement('p:grpSpPr');transform=OxmlElement('a:xfrm');b=e['box']
        for tag,values in [('a:off',{'x':b['left'],'y':b['top']}),('a:ext',{'cx':b['width'],'cy':b['height']}),('a:chOff',{'x':b['left'],'y':b['top']}),('a:chExt',{'cx':b['width'],'cy':b['height']})]:
            item=OxmlElement(tag)
            for key,value in values.items():item.set(key,str(round(value*EMU)))
            transform.append(item)
        prop.append(transform);node.append(prop)
        for item in objects:node.append(item)
    next_id=slide.shapes._next_shape_id
    for item in node.iter(f'{{{P}}}cNvPr'):
        item.set('id',str(next_id));next_id+=1
    node.find(f'.//{{{P}}}cNvPr').set('name',f"pppt:{slide._pppt_id}:{e['id']}")
    slide.shapes._spTree.insert_element_before(node,'p:extLst')
    return slide.shapes[-1],events


def populate(slide, elements, scene, base, scratch, converter):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    objects={};trace=[];pending=[]
    for e in sorted(elements,key=lambda v:v.get('layer',0)):
        b=e['box'];args=[round(b[k]*EMU) for k in ('left','top','width','height')]
        kind=e['kind'];details=[]
        if kind=='text':
            obj=slide.shapes.add_textbox(*args);put_text(obj,e)
        elif kind=='shape':
            obj,details=geometry_shape(slide,e,scene,base,scratch,converter)
            if e.get('visual_style',{}).get('dash'):
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                if not getattr(obj,'line',None):raise ValueError('分组矢量的虚线请写入各子形状素材')
                obj.line.dash_style={'solid':MSO_LINE_DASH_STYLE.SOLID,'dash':MSO_LINE_DASH_STYLE.DASH,'dot':MSO_LINE_DASH_STYLE.ROUND_DOT}[e['visual_style']['dash']]
            if e.get('label'):
                # A labelled node is one editable shape, not an image plus fake text.
                if not getattr(obj,'has_text_frame',False):raise ValueError('分组矢量节点的 label 请独立为 text 对象')
                put_text(obj,dict(e,text=e['label']))
        elif kind=='image':obj=slide.shapes.add_picture(str(asset(base,e['asset_path'])),*args)
        elif kind=='connector':
            start=e.get('start_point') or {'x':b['left'],'y':b['top']}
            end=e.get('end_point') or {'x':b['left']+b['width'],'y':b['top']+b['height']}
            obj=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,round(start['x']*EMU),round(start['y']*EMU),round(end['x']*EMU),round(end['y']*EMU))
            style=e.get('visual_style',{})
            obj.line.color.rgb=RGBColor.from_string(color(style.get('stroke')).lstrip('#'))
            obj.line.width=Pt(style.get('stroke_width_pt',1))
            if style.get('dash'):
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                obj.line.dash_style={'solid':MSO_LINE_DASH_STYLE.SOLID,'dash':MSO_LINE_DASH_STYLE.DASH,'dot':MSO_LINE_DASH_STYLE.ROUND_DOT}[style['dash']]
            from pptx.oxml.xmlchemy import OxmlElement
            line=obj._element.find(f'{{{P}}}spPr/{{{A}}}ln')
            for field,tag in [('start_arrow','headEnd'),('end_arrow','tailEnd')]:
                if style.get(field):
                    arrow=OxmlElement('a:'+tag);arrow.set('type',style[field]);line.append(arrow)
            pending.append((obj,e))
        elif kind=='table':
            rows=e['rows'];obj=slide.shapes.add_table(len(rows),len(rows[0]),*args)
            for ri,row in enumerate(rows):
                for ci,value in enumerate(row):
                    put_text(obj.table.cell(ri,ci),{'text':str(value),'font_size_pt':e.get('font_size_pt',16),'bold':ri==0})
        elif kind=='chart':
            types={'bar':XL_CHART_TYPE.BAR_CLUSTERED,'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'line':XL_CHART_TYPE.LINE,'pie':XL_CHART_TYPE.PIE}
            if e['chart_type'] not in types:raise ValueError('首批导出图表类型未支持: '+e['chart_type'])
            data=CategoryChartData();data.categories=e['categories']
            for series in e['series']:data.add_series(series['name'],series['values'])
            obj=slide.shapes.add_chart(types[e['chart_type']],*args,data)
            # Keep axis identities positive for readers that parse axId as uint.
            cns='http://schemas.openxmlformats.org/drawingml/2006/chart'
            for tag in ('axId','crossAx'):
                for axis in obj.chart._chartSpace.iter('{'+cns+'}'+tag):
                    axis.set('val',str(abs(int(axis.get('val')))))
            obj.chart.font.name=FONTS['en']
            obj.chart.font.size=Pt(e.get('font_size_pt',16))
            obj.chart.font.color.rgb=RGBColor.from_string(color(e.get('color')).lstrip('#'))
            for series,source in zip(obj.chart.series,e['series']):
                if source.get('color'):
                    rgb=RGBColor.from_string(color(source['color']).lstrip('#'))
                    series.format.fill.solid();series.format.fill.fore_color.rgb=rgb
                    series.format.line.color.rgb=rgb
            if e.get('unit') and e['chart_type']!='pie':
                obj.chart.value_axis.has_title=True
                put_text(obj.chart.value_axis.axis_title,{'text':e['unit'],'font_size_pt':e.get('font_size_pt',16)})
            # Native chart text inherits explicit family roles from its text body.
            for body in obj.chart._chartSpace.iter(f'{{{A}}}defRPr'):
                from pptx.oxml.xmlchemy import OxmlElement
                for tag,family in [('latin',FONTS['en']),('ea',FONTS['zh'])]:
                    node=OxmlElement('a:'+tag);node.set('typeface',family);body.append(node)
        else:raise ValueError('不支持的对象类型: '+str(kind))
        obj.name=f"pppt:{slide._pppt_id}:{e['id']}";objects[e['id']]=obj
        trace.append({'id':e['id'],'kind':kind,'name':obj.name,'conversion':'ppt-master-svg' if kind=='shape' else 'native-package-object','upstream_trace':details})
    for connector,e in pending:
        source,target=objects[e['from']],objects[e['to']]
        right=source.left+source.width/2<=target.left+target.width/2
        connector.begin_connect(source,3 if right else 1)
        connector.end_connect(target,1 if right else 3)
    return trace


def export_scene(scene, base, output, template_output=None, master_only=False):
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    errors=validate_scene(scene,base)
    if errors:raise ValueError('；'.join(errors))
    output=Path(output)
    if output.exists():raise ValueError('不覆盖已有 PPTX: '+str(output))
    converter=load_converter()
    presentation=Presentation()
    presentation.slide_width=round(scene['canvas']['width']*EMU);presentation.slide_height=round(scene['canvas']['height']*EMU)
    report={'scene_sha256':digest(scene),'engine':'PPT Master 604ecdec native SVG converter + python-pptx packaging','master':[],'slides':[]}
    with tempfile.TemporaryDirectory(prefix='.personal-ppt-native-',dir=Path(base).resolve()) as directory:
        scratch=Path(directory)
        if scene.get('master',{}).get('elements'):
            holder=presentation.slides.add_slide(presentation.slide_layouts[6]);holder._pppt_id='master'
            report['master']=populate(holder,scene['master']['elements'],scene,base,scratch,converter)
            master=presentation.slide_master
            for node in list(holder.shapes._spTree):
                if node.tag.rsplit('}',1)[-1] not in {'sp','pic','grpSp','cxnSp'}:continue
                node=copy.deepcopy(node)
                next_id=max((int(n.get('id','1')) for n in master.shapes._spTree.iter(f'{{{P}}}cNvPr')),default=1)+1
                for nv in node.iter(f'{{{P}}}cNvPr'):
                    nv.set('id',str(next_id));next_id+=1
                for item in node.iter():
                    for attribute in (f'{{{R}}}embed',f'{{{R}}}link'):
                        rid=item.get(attribute)
                        if rid:
                            target=holder.part.related_part(rid)
                            item.set(attribute,master.part.relate_to(target,RT.IMAGE))
                master.shapes._spTree.insert_element_before(node,'p:extLst')
            sid=presentation.slides._sldIdLst[-1];presentation.part.drop_rel(sid.rId);presentation.slides._sldIdLst.remove(sid)
        if template_output:
            template_output=Path(template_output)
            if template_output.exists():raise ValueError('不覆盖已有母版 PPTX')
            template_output.parent.mkdir(parents=True,exist_ok=True);presentation.save(template_output)
        for page in ([] if master_only else scene['slides']):
            slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide._pppt_id=page['id']
            from pptx.dml.color import RGBColor
            slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(color(page.get('background'),'#FFFFFF').lstrip('#'))
            report['slides'].append({'id':page['id'],'objects':populate(slide,page['elements'],scene,base,scratch,converter)})
        output.parent.mkdir(parents=True,exist_ok=True)
        presentation.save(output)
    report['pptx_sha256']=hashlib.sha256(output.read_bytes()).hexdigest()
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command',choices=['compile','layer-jobs','master','export'])
    parser.add_argument('--spec',required=True);parser.add_argument('--out',required=True)
    parser.add_argument('--stage',choices=['sample','production'],default='production')
    parser.add_argument('--template-out');parser.add_argument('--report')
    args=parser.parse_args()
    path=Path(args.spec).resolve();spec=load(path)
    import ppt_tools
    errors,_=ppt_tools.preflight(spec,path.parent,'layout' if args.command!='export' else args.stage)
    if errors:raise ValueError('；'.join(errors))
    if args.command=='layer-jobs':
        write(args.out,layer_jobs(spec,path.parent))
        print(json.dumps({'status':'pass','output':str(Path(args.out).resolve())},ensure_ascii=False));return
    scene=compile_scene(spec,path.parent)
    if args.command=='compile':write(args.out,scene)
    else:
        report=export_scene(scene,path.parent,args.out,args.template_out,master_only=args.command=='master')
        write(args.report or str(args.out)+'.report.json',report)
    print(json.dumps({'status':'pass','output':str(Path(args.out).resolve())},ensure_ascii=False))


if __name__=='__main__':
    try:main()
    except (OSError,ValueError,KeyError,ImportError,ET.ParseError) as error:
        print(json.dumps({'status':'fail','error':str(error)},ensure_ascii=False));raise SystemExit(2)
