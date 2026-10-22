"""Regression fixtures only. No fixture image is an imagegen result or a user approval."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/personal-ppt/scripts'))
import ppt_tools as ppt
import editorial_tools as editorial
import fidelity_tools as fidelity
import reconstruction_tools as reconstruction
import workflow_policy as policy
import design_contract
from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Regression(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='personal-ppt-v2-test-')
        self.base=Path(self.temp.name)
        for folder in ('design','build','prompts'):
            (self.base/folder).mkdir()
        image=self.base/'design/reference.png'
        Image.new('RGB',(1280,720),'white').save(image)
        self.image_sha=sha(image)
        self.spec={
            'schema_version':1,'workflow_version':'2.0','title':'Quarterly report',
            'audience':'Managers','purpose':'Decide the next step','language':'en',
            'mode':'per-slide','interaction':'autonomous','density':'speaking',
            'task_profile':{'scenario':'report','objective':'Choose a plan','delivery_context':'live'},
            'slide_count':1,'canvas':{'width':1280,'height':720},
            'style':{'name':'plain','palette':['#000000','#FFFFFF'],'fonts':['Microsoft YaHei','Times New Roman'],'visual_direction':'Clear evidence and generous space'},
            'sources':[],
            'master':{'id':'main','reference_image':'design/reference.png','reference_size':{'width':1280,'height':720},'image_sha256':self.image_sha,
                'approval':{'selected_by':'agent','confirmed':False,'selection_reason':'Suits a management report','user_selection':'','reference_sha256':self.image_sha},
                'tokens':{'fonts':{'zh':'Microsoft YaHei','en':'Times New Roman'},'text_styles':{'title':{'font_size_pt':28}}},
                'ppt_template_path':'build/sample.pptx'},
            'slides':[{'id':'slide-01','title':'Quarterly report','purpose':'INTERNAL PURPOSE SENTINEL',
                'layout':'One heading with space','layout_family':'typographic','treatment':'light','page_rhythm':'anchor',
                'design_reference':'design/reference.png','reference_size':{'width':1280,'height':720},
                'image_frame':{'left':0,'top':0,'width':1280,'height':720},
                'design_features':['Heading leads','Dark text','Clear margins'],
                'design_generation':{'master_sha256':self.image_sha},
                'elements':[{'id':'heading','kind':'text','editable':True,'text':'Delivery improved',
                    'text_channel':'slide-body','content_role':'research','master_slot':'title','font_size_pt':28,
                    'box':{'left':100,'top':100,'width':600,'height':100},
                    'reference_box':{'left':100,'top':100,'width':600,'height':100},'layer':0,
                    'runs':[{'text':'Delivery improved','font_family':'Times New Roman','font_size_pt':28,'color':'#000000','bold':True,'italic':False}],
                    'paragraph':{},'line_breaks':[]}]}]}
        self.spec['slides'][0]['reference_inventory']=[{'id':'heading','reference_box':self.spec['slides'][0]['elements'][0]['box'],'route':'native','element_ids':['heading']}]
        self.write_review()

    def tearDown(self):
        self.temp.cleanup()

    def write_review(self):
        report={'status':'pass','content_sha256':editorial.digest(self.spec),'observations':['Fixture copy reviewed'],
                **{k:True for k in ('channels_reviewed','facts_preserved','claim_strength_reviewed','redundancy_reviewed','semantic_wrap_reviewed')}}
        (self.base/'build/content-review.json').write_text(json.dumps(report),encoding='utf-8')

    def make_sample(self):
        presentation=Presentation()
        presentation.slide_width=Emu(1280*9525)
        presentation.slide_height=Emu(720*9525)
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        box=slide.shapes.add_textbox(Emu(100*9525),Emu(100*9525),Emu(600*9525),Emu(100*9525))
        box.name='pppt:slide-01:heading'
        run=box.text_frame.paragraphs[0].add_run()
        run.text='Delivery improved'
        run.font.name='Times New Roman'
        run.font.size=Pt(28)
        run.font.bold=True
        run.font.italic=False
        run.font.color.rgb=RGBColor(0,0,0)
        presentation.save(self.base/'build/sample.pptx')
        measurement={'origin':'visual-reference-measurement','slides':[{'id':'slide-01',
            'reference_sha256':self.image_sha,'reference_size':self.spec['slides'][0]['reference_size'],
            'image_frame':self.spec['slides'][0]['image_frame'],'observations':['Synthetic unit fixture'],
            'elements':copy.deepcopy(self.spec['slides'][0]['elements'])}]}
        (self.base/'design/reference-measurements.json').write_text(json.dumps(measurement),encoding='utf-8')
        (self.base/'build/comparison.html').write_text('<p>Unit fixture, not a rendered presentation review.</p>',encoding='utf-8')
        def record(path):return {'path':path,'sha256':sha(self.base/path)}
        receipt={'reviewed':True,'reviewer':'agent','user_reply':'','slide_id':'slide-01','complexity':['text-layout'],
            'displayed':True,'observations':['Synthetic file binding fixture, no production approval'],
            'blockers':{k:False for k in fidelity.BLOCKERS},'differences':[],
            'pptx':record('build/sample.pptx'),'reference':record('design/reference.png'),
            'render':record('design/reference.png'),'comparison':record('build/comparison.html'),
            'measurements':record('design/reference-measurements.json')}
        (self.base/'build/sample-approval.json').write_text(json.dumps(receipt),encoding='utf-8')
        return receipt

    def test_all_scenarios_plan_without_research_fields(self):
        for scenario in ('report','business','teaching','product','research','personal','general'):
            with self.subTest(scenario=scenario):
                self.spec['task_profile']['scenario']=scenario
                self.assertEqual([],ppt.preflight(self.spec,self.base,'planning')[0])

    def test_auto_selection_is_not_user_approval(self):
        self.assertEqual([],policy.check_master_selection(self.spec,self.image_sha))
        self.spec['master']['approval']['confirmed']=True
        self.assertTrue(policy.check_master_selection(self.spec,self.image_sha))

    def test_review_requires_user_selection(self):
        self.spec['interaction']='review'
        self.assertTrue(policy.check_master_selection(self.spec,self.image_sha))
        self.spec['master']['approval'].update(selected_by='user',confirmed=True,user_selection='Choose the first')
        self.assertEqual([],policy.check_master_selection(self.spec,self.image_sha))

    def test_legacy_selection_stays_strict(self):
        self.spec['workflow_version']='1.7'
        self.assertTrue(policy.check_master_selection(self.spec,self.image_sha))

    def test_master_hash_change_rejected(self):
        self.assertTrue(policy.check_master_selection(self.spec,'different-image'))

    def test_no_new_production_route(self):
        self.spec['mode']='representative'
        self.assertTrue(policy.check_task(self.spec))

    def test_v2_inference_not_forced_to_chinese_keywords(self):
        self.spec['slides'][0]['elements'][0].update(text='The mechanism remains a hypothesis',claim_type='inference')
        self.assertEqual([],editorial.audit(self.spec)[0])
        self.spec['workflow_version']='1.7'
        self.assertTrue(editorial.audit(self.spec)[0])

    def test_clean_copy_has_no_generic_per_page_warning(self):
        self.assertEqual(([],[]),editorial.audit(self.spec))

    def test_v2_style_warning_does_not_block(self):
        self.spec['slides'][0]['elements'][0]['text']='当前解释：交付时间变短'
        errors,warnings=editorial.audit(self.spec)
        self.assertEqual([],errors)
        self.assertTrue(warnings)

    def test_production_copy_still_blocked(self):
        self.spec['slides'][0]['elements'][0]['text']='制作记录：定义待核对'
        self.assertTrue(editorial.audit(self.spec)[0])

    def test_production_copy_in_table_blocked(self):
        self.spec['slides'][0]['elements'].append({'id':'table','kind':'table','rows':[['制作记录','待核对']]})
        self.assertTrue(editorial.audit(self.spec)[0])

    def test_reader_condition_exception_preserved(self):
        self.spec['slides'][0]['elements'][0].update(text='定义待核对',content_role='reader-condition',source_ids=['source-01'],editorial_exception={'reason':'Explicitly requested by author'})
        self.assertEqual([],editorial.audit(self.spec)[0])

    def test_prompt_excludes_internal_records(self):
        self.spec['speaker_notes']=['SPEAKER ONLY SENTINEL']
        self.spec['production_records']=['PRODUCTION ONLY SENTINEL']
        self.spec['slides'][0]['production_records']=['PAGE RECORD SENTINEL']
        jobs=ppt.prompt_jobs(self.spec,self.base,[])
        text=json.dumps(jobs,ensure_ascii=False)
        for marker in ('SPEAKER ONLY SENTINEL','PRODUCTION ONLY SENTINEL','PAGE RECORD SENTINEL','INTERNAL PURPOSE SENTINEL'):
            self.assertNotIn(marker,text)
        self.assertIn('Delivery improved',text)
        self.assertIn('Microsoft YaHei',text)
        self.assertIn('Times New Roman',text)

    def test_prompt_preserves_table_and_chart_values(self):
        self.spec['slides'][0]['elements'].extend([
            {'id':'table','kind':'table','rows':[['Period','Count'],['Q1',32]]},
            {'id':'chart','kind':'chart','categories':['Q1'],'series':[{'name':'Count','values':[32]}],'unit':'orders','chart_type':'bar'}])
        self.write_review()
        items=ppt.prompt_jobs(self.spec,self.base,[])[0]['visible_content']
        self.assertEqual(32,items[1]['rows'][1][1])
        self.assertEqual([32],items[2]['series'][0]['values'])

    def test_stale_copy_review_blocks_prompt(self):
        self.spec['slides'][0]['elements'][0]['text']='Changed after review'
        with self.assertRaises(ValueError):ppt.prompt_jobs(self.spec,self.base,[])

    def test_task_context_change_invalidates_copy_review(self):
        self.spec['task_profile']['delivery_context']='standalone'
        self.assertTrue(editorial.check(self.spec,self.base))

    def test_internal_approval_cannot_be_marked_as_user_approved(self):
        receipt=self.make_sample()
        receipt['approved']=True
        (self.base/'build/sample-approval.json').write_text(json.dumps(receipt),encoding='utf-8')
        self.assertTrue(fidelity.check_sample(self.spec,self.base))

    def test_unknown_slide_id_rejected(self):
        with self.assertRaises(ValueError):ppt.prompt_jobs(self.spec,self.base,['nonexistent'])

    def test_imagegen_receipt_and_real_images_still_required(self):
        self.assertTrue(ppt.image_design_gate(self.spec,self.base))

    def test_fonts_remain_enforced(self):
        self.spec['master']['tokens']['fonts']['zh']='Other font'
        self.assertTrue(any('微软雅黑' in e for e in reconstruction.check_contract(self.spec,self.base,'pagination')))

    def test_auto_sample_uses_actual_native_pptx(self):
        self.make_sample()
        self.assertEqual([],fidelity.check_sample(self.spec,self.base))

    def test_review_sample_still_requires_real_user_reply(self):
        self.make_sample()
        self.spec['interaction']='review'
        self.assertTrue(any('用户验收' in e for e in fidelity.check_sample(self.spec,self.base)))

    def test_sample_without_receipt_does_not_pass(self):
        self.assertTrue(fidelity.check_sample(self.spec,self.base))

    def test_auto_sample_rejects_forged_user_reply(self):
        receipt=self.make_sample()
        receipt['user_reply']='fabricated'
        (self.base/'build/sample-approval.json').write_text(json.dumps(receipt),encoding='utf-8')
        self.assertTrue(fidelity.check_sample(self.spec,self.base))

    def test_stale_sample_pptx_is_rejected(self):
        self.make_sample()
        with (self.base/'build/sample.pptx').open('ab') as f:f.write(b'changed')
        self.assertTrue(fidelity.check_sample(self.spec,self.base))

    def test_lock_tracks_task_and_interaction(self):
        snapshot=design_contract.snapshot(self.spec)
        self.spec['interaction']='review'
        self.assertNotEqual(snapshot,design_contract.snapshot(self.spec))
        self.spec['workflow_version']='1.7'
        self.assertNotIn('task_profile',design_contract.snapshot(self.spec))

    def test_init_sets_v2_and_same_fonts(self):
        destination=self.base/'new-project'
        ppt.initialize(destination)
        spec=json.loads((destination/'deck-spec.json').read_text(encoding='utf-8'))
        self.assertEqual('2.1',spec['workflow_version'])
        self.assertEqual('per-slide',spec['mode'])
        self.assertEqual(['Microsoft YaHei','Times New Roman'],spec['style']['fonts'])


if __name__=='__main__':
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(Regression))
    report={'tests':result.testsRun,'failures':len(result.failures),'errors':len(result.errors),'success':result.wasSuccessful(),
            'scope':'Synthetic regression fixtures and real OOXML specimen; no production imagegen or visual end-to-end claim.'}
    print(json.dumps(report,indent=2))
    raise SystemExit(not result.wasSuccessful())
