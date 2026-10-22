"""Synthetic export fixtures. These are not imagegen receipts or user approvals."""
import copy
import json
import sys
import tempfile
import unittest
import subprocess
import os
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/personal-ppt/scripts'))
import native_bridge as bridge
import ppt_tools
from PIL import Image
from pptx import Presentation


def fixture(base):
    # Transparent fixture pixels test independent layer transport, not raster editing.
    Image.new('RGBA',(400,200),(37,148,128,120)).save(base/'layer.png')
    (base/'mark.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><rect x="0" y="0" width="100" height="100" fill="#285A89"/><circle cx="50" cy="50" r="28" fill="#FFFFFF"/></svg>',encoding='utf-8')
    def box(x,y,w,h):return dict(left=x,top=y,width=w,height=h)
    elements=[
        {'id':'title','kind':'text','text':'原生转换验证 Native objects','box':box(64,44,1120,70),'font_size_pt':30,'bold':True,'color':'#16334B','layer':2},
        {'id':'subtitle','kind':'text','text':'技术验证，全部数值均为演示数据。','box':box(66,116,1000,35),'font_size_pt':16,'color':'#526579','layer':2},
        {'id':'panel','kind':'shape','geometry':'roundRect','box':box(64,176,500,400),'visual_style':{'fill':'#F0F5FA','stroke':'#D6E0E9','stroke_width_pt':1},'layer':0},
        {'id':'value','kind':'text','text':'68%','box':box(90,204,220,80),'font_size_pt':46,'bold':True,'color':'#008773','layer':2},
        {'id':'description','kind':'text','text':'独立图层、原生路径与混排文字','box':box(90,304,440,40),'font_size_pt':20,'layer':2},
        {'id':'picture','kind':'image','role':'illustration','asset_path':'layer.png','box':box(100,364,360,180),'layer':1},
        {'id':'step-a','kind':'shape','geometry':'preset:chevron','box':box(640,200,190,84),'visual_style':{'fill':'#285A89','stroke':'none'},'layer':1},
        {'id':'step-b','kind':'shape','geometry':'preset:hexagon','box':box(954,200,190,84),'visual_style':{'fill':'#008773','stroke':'none'},'layer':1},
        {'id':'label-a','kind':'text','text':'设计图','box':box(665,222,125,32),'font_size_pt':20,'bold':True,'color':'#FFFFFF','alignment':'center','layer':3},
        {'id':'label-b','kind':'text','text':'原生对象','box':box(975,222,144,32),'font_size_pt':20,'bold':True,'color':'#FFFFFF','alignment':'center','layer':3},
        {'id':'link','kind':'connector','from':'step-a','to':'step-b','box':box(830,242,124,0),'visual_style':{'stroke':'#526579','stroke_width_pt':1.5,'end_arrow':'triangle'},'layer':1},
        {'id':'mark','kind':'shape','geometry':'svg','svg_asset_path':'mark.svg','box':box(650,364,100,100),'layer':1},
        {'id':'path','kind':'shape','geometry':'freeform','svg_path':'M 800 460 L 800 360 Q 890 330 940 400 L 940 460 Z','box':box(800,340,140,120),'visual_style':{'fill':'#F2BC55','stroke':'none'},'layer':1},
        {'id':'caption','kind':'text','text':'矢量分组与自由曲线可分别编辑','box':box(650,488,500,35),'font_size_pt':18,'color':'#526579','layer':3}]
    master=[{'id':'footer','kind':'text','text':'personal-ppt 2.1  ·  技术验证','box':box(64,666,800,26),'font_size_pt':12,'color':'#6B7C8B'}]
    spec={'mode':'per-slide','canvas':{'width':1280,'height':720},'master':{'tokens':{'fonts':bridge.FONTS},'elements':master},
          'slides':[{'id':'demo-01','background':'#FFFFFF','elements':elements},
                    {'id':'demo-02','background':'#FFFFFF','elements':[
                        {'id':'heading','kind':'text','text':'原数据继续可编辑','box':box(64,44,1100,60),'font_size_pt':30,'bold':True},
                        {'id':'table','kind':'table','rows':[['阶段','数量'],['A',32],['B',48]],'box':box(64,200,440,260),'font_size_pt':20},
                        {'id':'chart','kind':'chart','chart_type':'column','categories':['A','B'],'series':[{'name':'演示数量','values':[32,48]}],'unit':'项','box':box(600,176,600,350)},
                        {'id':'note','kind':'text','text':'演示数据，不代表实际项目结果。','box':box(64,570,900,35),'font_size_pt':18,'color':'#526579'}]}]}
    return spec


class NativeBridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='personal-ppt-native-test-')
        self.base=Path(self.temp.name)
        self.spec=fixture(self.base)
        self.scene=bridge.compile_scene(self.spec,self.base)
    def tearDown(self):self.temp.cleanup()
    def test_master_only_has_no_pages(self):
        out=self.base/'template-only.pptx'
        bridge.export_scene(self.scene,self.base,out,master_only=True)
        deck=Presentation(out)
        self.assertEqual(0,len(deck.slides))
        self.assertTrue(any(s.name=='pppt:master:footer' for s in deck.slide_master.shapes))
    def test_master_style_is_compiled_once(self):
        self.spec['master']['tokens']['text_styles']={'heading':{'font_size_pt':27,'bold':True,'alignment':'center'}}
        e=self.spec['slides'][0]['elements'][0]
        e.pop('font_size_pt');e.pop('bold');e['master_slot']='heading'
        compiled=bridge.compile_scene(self.spec,self.base)['slides'][0]['elements'][0]
        self.assertEqual(27,compiled['font_size_pt']);self.assertTrue(compiled['bold'])
        self.assertEqual('center',compiled['alignment'])
    def test_new_policy_keeps_per_slide_and_review_requirements(self):
        import workflow_policy
        spec={'workflow_version':'2.1','interaction':'autonomous'}
        self.assertFalse(workflow_policy.needs_user_review(spec))
        spec['interaction']='review'
        self.assertTrue(workflow_policy.needs_user_review(spec))
    def test_public_compile_and_master_commands(self):
        from test_workflow_v2 import Regression
        case=Regression();case.setUp()
        try:
            case.spec['workflow_version']='2.1'
            spec=case.base/'deck-spec.json';spec.write_text(json.dumps(case.spec),encoding='utf-8')
            for command,suffix in [('compile','json'),('master','pptx')]:
                output=case.base/(command+'.'+suffix)
                result=subprocess.run([sys.executable,str(ROOT/'skills/personal-ppt/scripts/native_bridge.py'),command,'--spec',str(spec),'--out',str(output)],capture_output=True,text=True,encoding='utf-8',env={**os.environ,'PYTHONIOENCODING':'utf-8'})
                self.assertEqual(0,result.returncode,result.stdout+result.stderr)
                self.assertTrue(output.is_file())
        finally:case.tearDown()
    def test_public_export_does_not_bypass_generation_gates(self):
        spec=self.base/'spec.json';spec.write_text(json.dumps(self.spec),encoding='utf-8')
        out=self.base/'blocked.pptx'
        result=subprocess.run([sys.executable,str(ROOT/'skills/personal-ppt/scripts/native_bridge.py'),'export','--spec',str(spec),'--out',str(out)],capture_output=True,text=True,encoding='utf-8',env={**os.environ,'PYTHONIOENCODING':'utf-8'})
        self.assertNotEqual(0,result.returncode);self.assertFalse(out.exists())
    def test_actual_pptx_inspector_sees_native_groups_and_positions(self):
        out,_=self.export()
        self.spec['slide_count']=2
        result=ppt_tools.inspect(self.spec,out)
        self.assertEqual('pass',result['status'],result['errors'])
        trace=result['slides'][0]['conversion_trace']
        mark=next(e for e in trace if e['element_id']=='mark')
        self.assertEqual('verified',mark['status'])
        self.assertEqual({'left':650.,'top':364.,'width':100.,'height':100.},mark['actual_box'])
    def export(self):
        out=self.base/'native.pptx'
        report=bridge.export_scene(self.scene,self.base,out,self.base/'master.pptx')
        return out,report
    def test_real_upstream_geometry_and_native_objects(self):
        out,report=self.export()
        deck=Presentation(out)
        self.assertEqual(2,len(deck.slides))
        self.assertTrue(any(v['conversion']=='ppt-master-svg' for v in report['slides'][0]['objects']))
        with zipfile.ZipFile(out) as z:
            xml=z.read('ppt/slides/slide1.xml').decode()
            self.assertIn('custGeom',xml)
            self.assertIn('prst="chevron"',xml)
            self.assertIn('pppt:demo-01:mark',xml)
            self.assertIn('grpSp',xml)
            self.assertIn('stCxn',xml)
            self.assertIn('endCxn',xml)
            self.assertTrue(any(n.endswith('.xlsx') for n in z.namelist()))
            self.assertIn('graphicFrame',z.read('ppt/slides/slide2.xml').decode())
    def test_real_master_and_shared_layout(self):
        out,_=self.export()
        with zipfile.ZipFile(out) as z:
            master=z.read('ppt/slideMasters/slideMaster1.xml').decode()
            self.assertIn('pppt:master:footer',master)
            for number in (1,2):
                self.assertNotIn('pppt:master:footer',z.read(f'ppt/slides/slide{number}.xml').decode())
                self.assertIn('slideLayout',z.read(f'ppt/slides/_rels/slide{number}.xml.rels').decode())
            tree=ET.fromstring(master)
            ids=[n.get('id') for n in tree.iter('{'+bridge.P+'}cNvPr')]
            self.assertEqual(len(ids),len(set(ids)))
    def test_mixed_fonts_and_bold_survive(self):
        out,_=self.export()
        with zipfile.ZipFile(out) as z:
            tree=ET.fromstring(z.read('ppt/slides/slide1.xml'))
            title=next(n for n in tree.iter('{'+bridge.P+'}sp') if n.find('.//{'+bridge.P+'}cNvPr').get('name')=='pppt:demo-01:title')
            families=set()
            for run in title.iter('{'+bridge.A+'}r'):
                prop=run.find('{'+bridge.A+'}rPr')
                self.assertEqual('1',prop.get('b'))
                families.add(prop.find('{'+bridge.A+'}latin').get('typeface'))
            self.assertEqual(set(bridge.FONTS.values()),families)
    def test_png_alpha_is_preserved(self):
        out,_=self.export()
        with zipfile.ZipFile(out) as z:
            data=[z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
            self.assertIn((self.base/'layer.png').read_bytes(),data)
    def test_ir_excludes_backstage_records(self):
        self.spec['speaker_notes']=['do not put in scene']
        self.spec['production_records']=['private pending task']
        self.spec['slides'][0]['purpose']='internal brief'
        scene=bridge.compile_scene(self.spec,self.base)
        text=json.dumps(scene)
        for value in ('do not put in scene','private pending task','internal brief'):self.assertNotIn(value,text)
    def test_changed_asset_does_not_silently_export(self):
        Image.new('RGBA',(400,200),(0,0,0,0)).save(self.base/'layer.png')
        self.assertTrue(bridge.validate_scene(self.scene,self.base))
    def test_no_image_stretch(self):
        self.scene['slides'][0]['elements'][5]['box']['height']=100
        self.assertTrue(bridge.validate_scene(self.scene,self.base))
    def test_svg_active_content_is_rejected(self):
        with self.assertRaises(ValueError):bridge.safe_vector(b'<svg><script>bad()</script></svg>')
    def test_registration_stays_on_same_canvas(self):
        members=[copy.deepcopy(self.scene['slides'][0]['elements'][5]) for _ in range(2)]
        for i,e in enumerate(members):e.update(id='layer-'+str(i),registration_group='scene')
        members[1]['box']['left']+=10
        self.scene['slides'][0]['elements']=members
        self.assertTrue(bridge.validate_scene(self.scene,self.base))
    def test_reference_layer_jobs_use_original_design(self):
        self.spec['slides'][0]['design_reference']='layer.png'
        self.spec['slides'][0]['elements'][5].update(layer_role='subject',asset_prompt='Keep the existing product')
        tasks=bridge.layer_jobs(self.spec,self.base)['jobs']
        self.assertTrue(tasks[0]['transparent_background'])
        self.assertEqual(str(self.base/'layer.png'),tasks[0]['canonical_reference']['path'])
    def test_evidence_is_not_regenerated(self):
        self.spec['slides'][0]['design_reference']='layer.png'
        self.spec['slides'][0]['elements'][5].update(layer_role='subject',role='evidence')
        self.assertEqual([],bridge.layer_jobs(self.spec,self.base)['jobs'])
    def test_original_file_is_not_overwritten(self):
        self.export()
        with self.assertRaises(ValueError):bridge.export_scene(self.scene,self.base,self.base/'native.pptx')


if __name__=='__main__':unittest.main(verbosity=2)
